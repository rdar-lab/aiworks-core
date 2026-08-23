"""Tests for pptx_generator — render_presentation_pptx renders valid PPTX from JSON."""

import json
from io import BytesIO
from pathlib import Path

from django.test import TestCase
from ..logic.pptx_generator import render_presentation_pptx
from ..logic.schema_validation_utils import validate_json_with_schema_file

PPT_SCHEMA_PATH = Path(__file__).parent.parent / "resources/schema/ppt_schema.json"


def _validate_against_schema(presentation_json):
    """Validate presentation JSON against ppt_schema.json using schema_validation_utils."""
    json_str = json.dumps(presentation_json)
    return validate_json_with_schema_file(json_str, str(PPT_SCHEMA_PATH))


class RenderPresentationPptxTests(TestCase):
    def test_renders_valid_pptx_with_title_slide(self):
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "title",
                    "title": "Q3 Sales Review",
                    "elements": [
                        {"type": "text", "position": {"x": 0.5, "y": 0.3, "width": 12, "height": 1}, "text": "Q3 Sales Review"}
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})

        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(prs.slide_width.inches, 13.33)
        self.assertEqual(prs.slide_height.inches, 7.5)
        self.assertEqual(len(prs.slides), 1)

        slide = prs.slides[0]
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        full_text = " ".join(texts)
        self.assertIn("Q3 Sales Review", full_text)

    def test_renders_multiple_slide_types(self):
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {"type": "title", "title": "Title Slide"},
                {"type": "section", "title": "Section Slide"},
                {
                    "type": "content",
                    "elements": [
                        {"type": "text", "position": {"x": 0.5, "y": 2.3, "width": 12, "height": 0.5}, "text": "Heading"},
                        {"type": "text", "position": {"x": 0.5, "y": 2.9, "width": 12, "height": 0.4}, "text": "Bullet one"},
                        {"type": "text", "position": {"x": 0.5, "y": 3.4, "width": 12, "height": 0.4}, "text": "Bullet two"},
                    ],
                },
                {"type": "closing", "title": "Thank You"},
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 4)

        all_texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    all_texts.append(shape.text_frame.text)
        full_text = " ".join(all_texts)
        self.assertIn("Title Slide", full_text)
        self.assertIn("Section Slide", full_text)
        self.assertIn("Thank You", full_text)
        self.assertIn("Heading", full_text)
        self.assertIn("Bullet one", full_text)
        self.assertIn("Bullet two", full_text)

    def test_renders_slide_with_elements(self):
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "text",
                            "position": {"x": 0.5, "y": 2.0, "width": 5, "height": 0.5},
                            "text": "Hello world",
                        },
                        {
                            "type": "rect",
                            "position": {"x": 1.0, "y": 2.6, "width": 3.0, "height": 1.5},
                            "fill": "#e94560",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)

    def test_approves_more_than_50_slides(self):
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {"type": "title", "title": f"Slide {i}"}
                for i in range(51)
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsNotNone(pptx_bytes)

    def test_uses_default_dimensions_and_colors(self):
        presentation_json = {
            "slides": [
                {"type": "title", "title": "Minimal"},
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)

    def test_string_color_slots_are_resolved(self):
        """Colors can be passed as hex strings (not just {hex: ...} dicts)."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary": "#1a365d",
                    "secondary": "#2c5282",
                    "accent": "#3182ce",
                    "background": "#ffffff",
                    "text": "#1a202c",
                },
                "fonts": {"heading": "Arial", "body": "Calibri"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "rect",
                            "position": {"x": 0, "y": 0, "width": 13.33, "height": 0.3},
                            "fill": "primary",
                        },
                        {
                            "type": "text",
                            "position": {"x": 0.5, "y": 0.5, "width": 12, "height": 0.5},
                            "text": "Title Text",
                        },
                        {
                            "type": "rect",
                            "position": {"x": 0, "y": 7.0, "width": 13.33, "height": 0.5},
                            "fill": "accent",
                        },
                        {
                            "type": "text",
                            "position": {"x": 0.5, "y": 7.05, "width": 12, "height": 0.4},
                            "text": "Footer text",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)
        texts = []
        for shape in prs.slides[0].shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        self.assertIn("Title Text", " ".join(texts))
        self.assertIn("Footer text", " ".join(texts))

    def test_renders_image_with_svg_data(self):
        """SVG data in image element is rasterised and embedded."""
        svg_content = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><rect width="100" height="100" fill="#e94560"/></svg>'
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "image",
                            "position": {"x": 1.0, "y": 1.0, "width": 3.0, "height": 3.0},
                            "svg_data": svg_content,
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)
        slide = prs.slides[0]
        has_picture = any(shape.shape_type == 13 for shape in slide.shapes)
        self.assertTrue(has_picture, "SVG should be rasterised and embedded as picture")

    def test_renders_elements_with_slot_field(self):
        """Elements can have optional slot field for organizational clarity."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "text",
                            "slot": "title",
                            "position": {"x": 0.5, "y": 0.3, "width": 12, "height": 0.8},
                            "text": "Slot Title",
                        },
                        {
                            "type": "text",
                            "slot": "body",
                            "position": {"x": 0.5, "y": 1.5, "width": 12, "height": 2.0},
                            "text": "Body content with slot",
                        },
                        {
                            "type": "rect",
                            "slot": "footer",
                            "position": {"x": 0, "y": 7.0, "width": 13.33, "height": 0.5},
                            "fill": "secondary",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)
        texts = []
        for shape in prs.slides[0].shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        full_text = " ".join(texts)
        self.assertIn("Slot Title", full_text)
        self.assertIn("Body content with slot", full_text)

    def test_image_element_prefers_svg_over_image_file(self):
        """When both svg_data and image_file_name are present, svg_data takes precedence."""
        svg_content = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><rect width="100" height="100" fill="#0000ff"/></svg>'
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "image",
                            "position": {"x": 1.0, "y": 1.0, "width": 3.0, "height": 3.0},
                            "svg_data": svg_content,
                            "image_file_name": "should_be_ignored.png",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)
        slide = prs.slides[0]
        has_picture = any(shape.shape_type == 13 for shape in slide.shapes)
        self.assertTrue(has_picture, "SVG should be used even when image_file_name is also present")

    def test_renders_table_element(self):
        """Table element renders correctly with rows and columns."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "table",
                            "position": {"x": 0.5, "y": 1.0, "width": 12, "height": 3.0},
                            "columns": [
                                {"header": "Name", "width": 3},
                                {"header": "Value", "width": 3},
                                {"header": "Status", "width": 3},
                            ],
                            "rows": [
                                ["Item A", "100", "Active"],
                                ["Item B", "200", "Pending"],
                                ["Item C", "300", "Complete"],
                            ],
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)

        slide = prs.slides[0]
        has_table = any(shape.shape_type == 19 for shape in slide.shapes)
        self.assertTrue(has_table, "Table should be rendered")

        table_shape = next(s for s in slide.shapes if s.shape_type == 19)
        table = table_shape.table
        self.assertEqual(len(table.rows), 4)
        self.assertEqual(len(table.columns), 3)
        self.assertEqual(table.cell(0, 0).text, "Name")
        self.assertEqual(table.cell(1, 0).text, "Item A")

    def test_renders_table_with_custom_table_config(self):
        """Table element respects template.table styling config."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#e94560",
                    "secondary": "#1a1a2e",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#000000",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
                "table": {
                    "header_fill": "#1a1a2e",
                    "header_font_color": "#ffffff",
                    "cell_fill_alt": "#f0f0f0",
                    "cell_font_color": "#333333",
                    "font_size": 16,
                    "font": {"family": "Arial", "bold": False}
                },
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "table",
                            "position": {"x": 0.5, "y": 1.0, "width": 12, "height": 3.0},
                            "columns": [
                                {"header": "Col A", "width": 6},
                                {"header": "Col B", "width": 6},
                            ],
                            "rows": [
                                ["Row 1 A", "Row 1 B"],
                                ["Row 2 A", "Row 2 B"],
                            ],
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)

        slide = prs.slides[0]
        has_table = any(shape.shape_type == 19 for shape in slide.shapes)
        self.assertTrue(has_table, "Table should be rendered with custom config")

    def test_text_element_with_font_color(self):
        """Text element can specify font_color."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "text",
                            "position": {"x": 0.5, "y": 1.0, "width": 5, "height": 0.5},
                            "text": "Red text",
                            "font_color": "#ff0000",
                        },
                        {
                            "type": "text",
                            "position": {"x": 0.5, "y": 1.8, "width": 5, "height": 0.5},
                            "text": "Primary color",
                            "font_color": "primary",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)

        texts = []
        for shape in prs.slides[0].shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        full_text = " ".join(texts)
        self.assertIn("Red text", full_text)
        self.assertIn("Primary color", full_text)

    def test_text_element_with_font_object(self):
        """Text element can specify font object with family, size, bold, italic."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "text",
                            "position": {"x": 0.5, "y": 1.0, "width": 6, "height": 0.6},
                            "text": "Styled text",
                            "font": {
                                "family": "Georgia",
                                "size": 24,
                                "bold": True,
                                "italic": True
                            },
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)

        texts = []
        for shape in prs.slides[0].shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        self.assertIn("Styled text", " ".join(texts))

    def test_renders_image_from_image_file_name(self):
        """Image element can load from image_file_name in images dict."""
        png_data = b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82'
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "image",
                            "position": {"x": 1.0, "y": 1.0, "width": 3.0, "height": 3.0},
                            "image_file_name": "chart.png",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {"chart.png": png_data})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)
        slide = prs.slides[0]
        has_picture = any(shape.shape_type == 13 for shape in slide.shapes)
        self.assertTrue(has_picture, "Image from file should be embedded")

    def test_align_property_on_text_element(self):
        """Text element supports align property (left, center, right)."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "text",
                            "position": {"x": 0.5, "y": 1.0, "width": 12, "height": 0.5},
                            "text": "Centered text",
                            "align": "center",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)

    def test_valid_presentation_json_passes_schema_validation(self):
        """Full presentation JSON passes schema validation."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary": "#1a1a2e",
                    "secondary": "#e94560",
                    "accent": "#0f3460",
                    "background": "#ffffff",
                    "text": "#333333",
                },
                "fonts": {"heading": {"family": "Arial"}, "body": {"family": "Georgia"}},
                "table": {
                    "header_fill": "#1a1a2e",
                    "header_font_color": "#ffffff"
                },
            },
            "slides": [
                {
                    "type": "title",
                    "title": "Test Presentation",
                    "elements": [
                        {"type": "text", "position": {"x": 0.5, "y": 0.3, "width": 12, "height": 1}, "text": "Title"},
                        {"type": "text", "slot": "body", "position": {"x": 0.5, "y": 1.5, "width": 12, "height": 2}, "text": "Body text", "font_color": "primary"},
                    ],
                },
                {
                    "type": "content",
                    "elements": [
                        {"type": "rect", "position": {"x": 0, "y": 0, "width": 13.33, "height": 0.5}, "fill": "primary"},
                        {"type": "image", "position": {"x": 1, "y": 1, "width": 4, "height": 4}, "image_file_name": "test.png"},
                        {"type": "table", "position": {"x": 6, "y": 1, "width": 6, "height": 3}, "columns": [{"header": "Col", "width": 6}], "rows": [["Cell"]]},
                    ],
                },
            ],
        }
        is_valid, error = _validate_against_schema(presentation_json)
        self.assertTrue(is_valid, error)
        pptx_bytes = render_presentation_pptx(presentation_json, {"test.png": b'\x89PNG\r\n\x1a\n'})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

    def test_schema_requires_position_on_elements(self):
        """Schema validation fails when elements missing required position."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {"primary": "#1a1a2e", "background": "#ffffff", "text": "#333333"},
                "fonts": {"heading": "Arial", "body": "Calibri"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {"type": "text", "text": "Missing position"},
                    ],
                },
            ],
        }
        is_valid, _ = _validate_against_schema(presentation_json)
        self.assertFalse(is_valid)

    def test_schema_requires_image_source(self):
        """Schema validation fails when image element has neither svg_data nor image_file_name."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {"primary": "#1a1a2e", "background": "#ffffff", "text": "#333333"},
                "fonts": {"heading": "Arial", "body": "Calibri"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {"type": "image", "position": {"x": 1, "y": 1, "width": 3, "height": 3}},
                    ],
                },
            ],
        }
        is_valid, _ = _validate_against_schema(presentation_json)
        self.assertFalse(is_valid)

    def test_rect_element_with_stroke(self):
        """Rect element can have stroke styling."""
        presentation_json = {
            "template": {
                "slide_width_inches": 13.33,
                "slide_height_inches": 7.5,
                "colors": {
                    "primary":   "#1a1a2e",
                    "secondary": "#e94560",
                    "accent":    "#0f3460",
                    "background": "#ffffff",
                    "text":      "#333333",
                },
                "fonts": {"heading": "Calibri", "body": "Arial"},
            },
            "slides": [
                {
                    "type": "content",
                    "elements": [
                        {
                            "type": "rect",
                            "position": {"x": 1.0, "y": 1.0, "width": 4.0, "height": 3.0},
                            "fill": "#e94560",
                            "stroke": "#1a1a2e",
                        },
                    ],
                },
            ],
        }

        pptx_bytes = render_presentation_pptx(presentation_json, {})
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertGreater(len(pptx_bytes), 0)

        from pptx import Presentation
        buf = BytesIO(pptx_bytes)
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)